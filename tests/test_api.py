import io
import os
import time
import zipfile

import pikepdf
import pytest
from fastapi.testclient import TestClient
from PIL import Image

from app.core.config import settings
from app.main import app
from app.services.file_service import cleanup_old_files
from app.services.pdf_service import parse_page_ranges


@pytest.fixture()
def client(tmp_path, monkeypatch):
    up, res = tmp_path / "uploads", tmp_path / "results"
    up.mkdir()
    res.mkdir()
    monkeypatch.setattr(settings, "UPLOAD_FOLDER", str(up))
    monkeypatch.setattr(settings, "RESULT_FOLDER", str(res))
    # Tes fungsional mengirim ratusan request dari satu "IP"; pengaman diuji khusus di test_ratelimit.py
    monkeypatch.setattr(settings, "RATE_LIMIT_ENABLED", False)
    with TestClient(app) as c:
        yield c


def png_bytes(size=(400, 300), color=(200, 30, 30), mode="RGB"):
    buf = io.BytesIO()
    Image.effect_noise(size, 80).convert(mode).save(buf, "PNG") if mode == "L" else \
        Image.new(mode, size, color if mode == "RGB" else color + (128,)).save(buf, "PNG")
    return buf.getvalue()


def noisy_jpeg(size=(1800, 1200)):
    img = Image.effect_noise(size, 90).convert("RGB")
    buf = io.BytesIO()
    img.save(buf, "JPEG", quality=95)
    return buf.getvalue()


def pdf_bytes(pages=3, image=None):
    """PDF sederhana dibuat lewat Pillow (tanpa reportlab)."""
    imgs = []
    for i in range(pages):
        if image is not None:
            imgs.append(image.copy())
        else:
            imgs.append(Image.new("RGB", (300, 400), (255 - i * 40, 255, 255)))
    buf = io.BytesIO()
    imgs[0].save(buf, "PDF", save_all=True, append_images=imgs[1:])
    return buf.getvalue()


def upload(client, name, data, mime="application/octet-stream"):
    r = client.post("/api/v1/upload", files=[("files", (name, data, mime))])
    assert r.status_code == 200, r.text
    return r.json()["data"][0]["saved_name"]


def process(client, action, names, **options):
    return client.post(
        "/api/v1/process",
        json={"action": action, "filenames": names, "original_names": [f"doc{i}" for i in range(len(names))], "options": options},
    )


def download(client, r):
    res = r.json()["results"][0]
    d = client.get(f"/api/v1/download/{res['processed_file']}", params={"name": res["download_name"]})
    assert d.status_code == 200
    return d


# ------------------------------------------------------------------ upload
def test_upload_rejects_fake_extension(client):
    r = client.post("/api/v1/upload", files=[("files", ("evil.pdf", b"MZ not a pdf", "application/pdf"))])
    assert r.status_code == 400


def test_upload_rejects_too_large(client, monkeypatch):
    monkeypatch.setattr(settings, "MAX_UPLOAD_MB", 0)
    r = client.post("/api/v1/upload", files=[("files", ("a.png", png_bytes(), "image/png"))])
    assert r.status_code == 400
    assert os.listdir(settings.UPLOAD_FOLDER) == []


def test_upload_reports_rejected_files(client):
    r = client.post("/api/v1/upload", files=[
        ("files", ("ok.png", png_bytes(), "image/png")),
        ("files", ("bad.exe", b"MZ....", "application/octet-stream")),
    ])
    body = r.json()
    assert len(body["data"]) == 1 and len(body["rejected"]) == 1


# ---------------------------------------------------------------- security
@pytest.mark.parametrize("bad", ["../../etc/passwd", "..\\secret.png", "a/b.png", "x"])
def test_process_rejects_path_traversal(client, bad):
    r = process(client, "compress", [bad])
    assert r.status_code == 400


def test_download_rejects_traversal(client, tmp_path):
    (tmp_path / "secret.txt").write_text("x")
    for bad in ["..%5Csecret.txt", "..%2Fsecret.txt", "secret.txt"]:
        assert client.get(f"/api/v1/download/{bad}").status_code == 404


def test_download_can_be_repeated(client):
    name = upload(client, "a.png", png_bytes())
    r = process(client, "convert", [name], format="jpg")
    download(client, r)
    download(client, r)   # sebelumnya file dihapus setelah download pertama


# ------------------------------------------------------------------- image
def test_compress_jpeg_smaller_and_named(client):
    data = noisy_jpeg()
    name = upload(client, "foto.jpg", data)
    r = process(client, "compress", [name], level="low")
    res = r.json()["results"][0]
    assert res["processed_size"] < len(data)
    assert res["download_name"] == "doc0_compressed.jpg"
    d = download(client, r)
    assert "doc0_compressed.jpg" in d.headers["content-disposition"]


def test_compress_target_kb(client):
    name = upload(client, "foto.jpg", noisy_jpeg())
    r = process(client, "compress", [name], target_kb=200)
    assert r.json()["results"][0]["processed_size"] <= 200 * 1024


def test_compress_png_output_extension_matches_content(client):
    name = upload(client, "scan.png", png_bytes(size=(1000, 1000)))
    r = process(client, "compress", [name])
    out = r.json()["results"][0]["processed_file"]
    with Image.open(os.path.join(settings.RESULT_FOLDER, out)) as img:
        assert out.endswith(".png") == (img.format == "PNG")


def test_resize_pixels_keep_ratio(client):
    name = upload(client, "a.png", png_bytes(size=(400, 200)))
    r = process(client, "resize", [name], mode="pixels", width=100)
    out = r.json()["results"][0]["processed_file"]
    with Image.open(os.path.join(settings.RESULT_FOLDER, out)) as img:
        assert img.size == (100, 50)


def test_resize_percent_and_validation(client):
    name = upload(client, "a.png", png_bytes(size=(400, 200)))
    r = process(client, "resize", [name], mode="percent", percent=50)
    out = r.json()["results"][0]["processed_file"]
    with Image.open(os.path.join(settings.RESULT_FOLDER, out)) as img:
        assert img.size == (200, 100)
    name = upload(client, "a.png", png_bytes())
    assert process(client, "resize", [name], mode="percent", percent="abc").status_code == 400


def test_convert_transparent_png_to_jpg_has_white_background(client):
    buf = io.BytesIO()
    Image.new("RGBA", (50, 50), (0, 0, 0, 0)).save(buf, "PNG")
    name = upload(client, "t.png", buf.getvalue())
    r = process(client, "convert", [name], format="jpg")
    out = r.json()["results"][0]["processed_file"]
    with Image.open(os.path.join(settings.RESULT_FOLDER, out)) as img:
        assert img.getpixel((25, 25))[0] > 240


def test_images_to_single_pdf_in_order(client):
    names = [upload(client, f"p{i}.png", png_bytes()) for i in range(3)]
    r = process(client, "pdf", names)
    out = r.json()["results"][0]["processed_file"]
    with pikepdf.open(os.path.join(settings.RESULT_FOLDER, out)) as pdf:
        assert len(pdf.pages) == 3


def test_image_tool_rejects_pdf(client):
    name = upload(client, "a.pdf", pdf_bytes(1))
    r = process(client, "compress", [name])
    assert r.status_code == 400


# --------------------------------------------------------------------- pdf
def test_compress_pdf_shrinks_image_heavy_pdf(client):
    big = Image.effect_noise((2400, 3200), 90).convert("RGB")
    data = pdf_bytes(2, image=big)
    name = upload(client, "scan.pdf", data)
    r = process(client, "compress-pdf", [name], level="low")
    res = r.json()["results"][0]
    assert res["processed_size"] < len(data) * 0.6
    with pikepdf.open(os.path.join(settings.RESULT_FOLDER, res["processed_file"])) as pdf:
        assert len(pdf.pages) == 2


def test_compress_pdf_target_not_reached_is_reported(client):
    name = upload(client, "a.pdf", pdf_bytes(1, image=Image.effect_noise((2000, 2000), 90).convert("RGB")))
    r = process(client, "compress-pdf", [name], level="high", target_kb=20)
    assert r.json()["results"][0]["note"] == "target_not_reached"


def test_compress_pdf_never_returns_bigger_file(client):
    data = pdf_bytes(1)
    name = upload(client, "tiny.pdf", data)
    r = process(client, "compress-pdf", [name])
    assert r.json()["results"][0]["processed_size"] <= len(data)


def test_merge_pdfs_and_image(client):
    a = upload(client, "a.pdf", pdf_bytes(2))
    b = upload(client, "b.pdf", pdf_bytes(3))
    c = upload(client, "c.png", png_bytes())
    r = process(client, "merge", [a, b, c])
    out = r.json()["results"][0]["processed_file"]
    with pikepdf.open(os.path.join(settings.RESULT_FOLDER, out)) as pdf:
        assert len(pdf.pages) == 6


def test_merge_requires_two_files(client):
    a = upload(client, "a.pdf", pdf_bytes(2))
    assert process(client, "merge", [a]).status_code == 400


def test_split_extract_remove_each(client):
    def pages_of(r):
        out = r.json()["results"][0]["processed_file"]
        with pikepdf.open(os.path.join(settings.RESULT_FOLDER, out)) as pdf:
            return len(pdf.pages)

    n = upload(client, "a.pdf", pdf_bytes(5))
    assert pages_of(process(client, "split", [n], mode="extract", ranges="1-2,5")) == 3
    n = upload(client, "a.pdf", pdf_bytes(5))
    assert pages_of(process(client, "split", [n], mode="remove", ranges="2-3")) == 3
    n = upload(client, "a.pdf", pdf_bytes(4))
    r = process(client, "split", [n], mode="each")
    z = zipfile.ZipFile(io.BytesIO(download(client, r).content))
    assert len(z.namelist()) == 4


def test_split_invalid_range_gives_clear_error(client):
    n = upload(client, "a.pdf", pdf_bytes(3))
    r = process(client, "split", [n], mode="extract", ranges="2-9")
    assert r.status_code == 400 and "3 halaman" in r.json()["detail"]


def test_rotate(client):
    n = upload(client, "a.pdf", pdf_bytes(2))
    r = process(client, "rotate", [n], angle=90, pages="1")
    out = r.json()["results"][0]["processed_file"]
    with pikepdf.open(os.path.join(settings.RESULT_FOLDER, out)) as pdf:
        assert int(pdf.pages[0].get("/Rotate", 0)) == 90
        assert int(pdf.pages[1].get("/Rotate", 0)) == 0


def test_protect_then_unlock(client):
    n = upload(client, "a.pdf", pdf_bytes(1))
    r = process(client, "protect", [n], password="rahasia123")
    protected = download(client, r).content
    with pytest.raises(pikepdf.PasswordError):
        pikepdf.open(io.BytesIO(protected))

    n2 = upload(client, "locked.pdf", protected)
    assert process(client, "unlock", [n2], password="salah").status_code == 400
    n3 = upload(client, "locked.pdf", protected)
    r = process(client, "unlock", [n3], password="rahasia123")
    pikepdf.open(io.BytesIO(download(client, r).content)).close()


def test_pdf_to_images(client):
    n = upload(client, "a.pdf", pdf_bytes(3))
    r = process(client, "pdf-to-image", [n], format="jpg", dpi=72)
    z = zipfile.ZipFile(io.BytesIO(download(client, r).content))
    assert len(z.namelist()) == 3
    n = upload(client, "one.pdf", pdf_bytes(1))
    r = process(client, "pdf-to-image", [n], format="png")
    assert r.json()["results"][0]["processed_file"].endswith(".png")


def test_partial_failure_is_reported_per_file(client):
    good = upload(client, "a.pdf", pdf_bytes(1))
    r = process(client, "compress-pdf", [good, "00000000-0000-0000-0000-000000000000.pdf"])
    body = r.json()
    assert len(body["results"]) == 1 and len(body["errors"]) == 1


def test_inputs_removed_after_processing(client):
    n = upload(client, "a.png", png_bytes())
    process(client, "convert", [n], format="jpg")
    assert os.listdir(settings.UPLOAD_FOLDER) == []


# ------------------------------------------------------------------- utils
def test_parse_page_ranges():
    assert parse_page_ranges("1-3,5,8-", 9) == [0, 1, 2, 4, 7, 8]
    assert parse_page_ranges("3,1,3", 5) == [2, 0]
    for bad in ["", "0", "4-2", "x", "1-99"]:
        with pytest.raises(ValueError):
            parse_page_ranges(bad, 5)


def test_cleanup_removes_only_old_files(client):
    old = os.path.join(settings.UPLOAD_FOLDER, "old.png")
    new = os.path.join(settings.RESULT_FOLDER, "new.png")
    for p in (old, new):
        open(p, "wb").write(b"x")
    past = time.time() - 3600
    os.utime(old, (past, past))
    assert cleanup_old_files(max_age_seconds=1800) == 1
    assert not os.path.exists(old) and os.path.exists(new)


# ------------------------------------------------------------ word <-> pdf
def docx_bytes(text="Halo dunia. Ini dokumen uji."):
    import docx
    d = docx.Document()
    d.add_heading("Judul Laporan", 1)
    for _ in range(3):
        d.add_paragraph(text)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def text_pdf_bytes():
    import pymupdf
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Laporan Praktikum Fisika", fontsize=18)
    page.insert_text((72, 110), "Percobaan ini membahas gerak lurus beraturan.", fontsize=11)
    return doc.tobytes()


def test_upload_accepts_real_docx_rejects_fake_zip(client):
    upload(client, "tugas.docx", docx_bytes())
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("readme.txt", "bukan docx")
    r = client.post("/api/v1/upload", files=[("files", ("fake.docx", buf.getvalue(), "application/zip"))])
    assert r.status_code == 400


def test_pdf_to_docx_keeps_text(client):
    import docx
    n = upload(client, "laporan.pdf", text_pdf_bytes())
    r = process(client, "pdf-to-docx", [n])
    assert r.json()["results"][0]["download_name"] == "doc0.docx"
    d = docx.Document(io.BytesIO(download(client, r).content))
    assert "Laporan Praktikum Fisika" in "\n".join(p.text for p in d.paragraphs)


def test_pdf_to_docx_rejects_protected_pdf(client):
    n = upload(client, "a.pdf", pdf_bytes(1))
    protected = download(client, process(client, "protect", [n], password="rahasia123")).content
    n2 = upload(client, "locked.pdf", protected)
    r = process(client, "pdf-to-docx", [n2])
    assert r.status_code == 400 and "password" in r.json()["detail"].lower()


def test_docx_tool_rejects_pdf_input(client):
    n = upload(client, "a.pdf", pdf_bytes(1))
    assert process(client, "docx-to-pdf", [n]).status_code == 400


@pytest.mark.skipif(not any(__import__("app.services.office_service", fromlist=["x"]).available_engines().values()),
                    reason="LibreOffice / Microsoft Word tidak tersedia")
def test_docx_to_pdf(client):
    import pymupdf
    n = upload(client, "tugas.docx", docx_bytes())
    r = process(client, "docx-to-pdf", [n])
    pdf = pymupdf.open(stream=download(client, r).content, filetype="pdf")
    assert pdf.page_count >= 1 and "Judul Laporan" in pdf[0].get_text()
