import io
import os

import numpy as np
import pymupdf
import pytest
from PIL import Image, ImageDraw, ImageFont

from app.core.config import settings
from app.services import ocr_service, office_service
from app.services.ocr_service import Word, group_paragraphs
from tests.test_api import client, download, png_bytes, process, upload  # noqa: F401

LINES = [
    "LAPORAN PRAKTIKUM FISIKA",
    "Percobaan tiga gerak lurus beraturan",
    "Mahasiswa mengukur kecepatan benda yang bergerak",
    "di atas lintasan lurus menggunakan stopwatch digital.",
]

needs_ocr = pytest.mark.skipif(not ocr_service.is_available(), reason="rapidocr tidak terpasang")


def scan_image(lines=LINES, size=(1240, 600)) -> Image.Image:
    img = Image.new("RGB", size, "white")
    d = ImageDraw.Draw(img)
    font = ImageFont.load_default(size=38)
    for i, line in enumerate(lines):
        d.text((60, 50 + i * 80), line, fill=(20, 20, 20), font=font)
    return img


def scanned_pdf(pages=1) -> bytes:
    """PDF hanya berisi gambar (tanpa teks digital), seperti hasil scan."""
    imgs = [scan_image() for _ in range(pages)]
    buf = io.BytesIO()
    imgs[0].save(buf, "PDF", save_all=True, append_images=imgs[1:], resolution=150)
    return buf.getvalue()


def image_bytes(img: Image.Image) -> bytes:
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


def digital_pdf() -> bytes:
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 100), "Dokumen ini sudah memiliki teks digital yang bisa dicari.", fontsize=12)
    return doc.tobytes()


# ---------------------------------------------------------------- OCR PDF
@needs_ocr
def test_ocr_pdf_makes_scan_searchable_without_changing_look(client):
    original = scanned_pdf()
    with pymupdf.open(stream=original, filetype="pdf") as doc:
        assert doc[0].get_text().strip() == ""
        before = np.frombuffer(doc[0].get_pixmap(dpi=60, alpha=False).samples, np.uint8).astype(int)

    n = upload(client, "scan.pdf", original)
    r = process(client, "ocr-pdf", [n])
    assert r.json()["results"][0]["download_name"] == "doc0_ocr.pdf"
    with pymupdf.open(stream=download(client, r).content, filetype="pdf") as doc:
        page = doc[0]
        text = page.get_text().upper()
        assert "LAPORAN" in text and "PRAKTIKUM" in text and "STOPWATCH" in text
        hits = page.search_for("PRAKTIKUM")
        assert hits and all(page.rect.contains(h) for h in hits)    # teks tepat berada di dalam halaman
        after = np.frombuffer(page.get_pixmap(dpi=60, alpha=False).samples, np.uint8).astype(int)
    assert np.abs(before - after).mean() < 1.0                      # teks tak terlihat: tampilan sama


@needs_ocr
def test_ocr_pdf_skips_pages_that_already_have_text(client):
    n = upload(client, "digital.pdf", digital_pdf())
    r = process(client, "ocr-pdf", [n])
    assert r.json()["results"][0]["note"] == "already_searchable"


@needs_ocr
def test_ocr_pdf_page_limit(client, monkeypatch):
    monkeypatch.setattr(settings, "MAX_OCR_PAGES", 1)
    n = upload(client, "scan.pdf", scanned_pdf(pages=2))
    r = process(client, "ocr-pdf", [n])
    assert r.status_code == 400 and "dibatasi 1 halaman" in r.json()["detail"]


# --------------------------------------------------------------- OCR teks
@needs_ocr
def test_ocr_text_image_to_docx(client):
    import docx
    n = upload(client, "foto buku.png", image_bytes(scan_image()))
    r = process(client, "ocr-text", [n], format="docx")
    assert r.json()["results"][0]["download_name"] == "doc0_teks.docx"
    doc = docx.Document(io.BytesIO(download(client, r).content))
    text = " ".join(p.text for p in doc.paragraphs).upper()
    assert "LAPORAN" in text and "LINTASAN" in text


@needs_ocr
def test_ocr_text_scanned_pdf_to_txt_with_page_separation(client):
    n = upload(client, "scan.pdf", scanned_pdf(pages=2))
    r = process(client, "ocr-text", [n], format="txt")
    content = download(client, r).content.decode("utf-8").upper()
    assert content.count("LAPORAN") == 2 and "STOPWATCH" in content


@needs_ocr
def test_ocr_text_uses_digital_text_without_ocr(client, monkeypatch):
    monkeypatch.setattr(ocr_service, "_recognize", lambda image: pytest.fail("tidak boleh OCR"))
    n = upload(client, "digital.pdf", digital_pdf())
    r = process(client, "ocr-text", [n], format="txt")
    assert "teks digital" in download(client, r).content.decode("utf-8")


@needs_ocr
def test_ocr_text_blank_image_gives_clear_error(client):
    n = upload(client, "kosong.png", png_bytes(size=(600, 400), color=(255, 255, 255)))
    r = process(client, "ocr-text", [n])
    assert r.status_code == 400 and "Tidak ada teks" in r.json()["detail"]


def test_ocr_text_rejects_unknown_format(client):
    n = upload(client, "a.png", png_bytes())
    assert process(client, "ocr-text", [n], format="rtf").status_code == 400


def test_ocr_unavailable_message(client, monkeypatch):
    monkeypatch.setattr(ocr_service, "is_available", lambda: False)
    n = upload(client, "a.png", png_bytes())
    r = process(client, "ocr-text", [n])
    assert r.status_code == 400 and "belum terpasang" in r.json()["detail"]


# ---------------------------------------------------------- paragraph logic
def box(text, x0, y0, x1, y1):
    return Word(text, x0, y0, x1, y1)


def test_group_paragraphs_merges_row_and_wrapped_lines():
    words = [
        box("Judul", 100, 10, 300, 40),                                     # baris pendek -> paragraf sendiri
        box("Baris", 100, 100, 300, 130), box("pertama", 320, 101, 700, 131),   # satu baris, dua kotak
        box("sambungan", 100, 140, 700, 170),                               # baris penuh -> menyambung
        box("akhir.", 100, 180, 300, 210),                                  # baris pendek -> menutup paragraf
        box("Paragraf", 100, 300, 700, 330),                                # jarak besar -> paragraf baru
    ]
    assert group_paragraphs(words) == ["Judul", "Baris pertama sambungan akhir.", "Paragraf"]


def test_group_paragraphs_empty():
    assert group_paragraphs([]) == []


# ------------------------------------------------------------------ PPTX
def pptx_bytes() -> bytes:
    from pptx import Presentation
    prs = Presentation()
    for title in ("Presentasi Skripsi", "Metodologi Penelitian"):
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = title
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_upload_accepts_real_pptx_and_keeps_extension(client):
    saved = upload(client, "slide.pptx", pptx_bytes())
    assert saved.endswith(".pptx") and os.path.exists(os.path.join(settings.UPLOAD_FOLDER, saved))


def test_upload_rejects_plain_zip(client):
    import zipfile
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("data.txt", "bukan office")
    r = client.post("/api/v1/upload", files=[("files", ("a.pptx", buf.getvalue(), "application/zip"))])
    assert r.status_code == 400


def test_pptx_tool_rejects_docx_and_pdf(client):
    n = upload(client, "a.pdf", digital_pdf())
    assert process(client, "pptx-to-pdf", [n]).status_code == 400


def test_pptx_to_pdf_unavailable_message(client, monkeypatch):
    monkeypatch.setattr(office_service, "available_engines",
                        lambda: {"libreoffice": False, "word": False, "powerpoint": False})
    n = upload(client, "slide.pptx", pptx_bytes())
    r = process(client, "pptx-to-pdf", [n])
    assert r.status_code == 400 and "LibreOffice" in r.json()["detail"]


@pytest.mark.skipif(not office_service.available_kinds()[".pptx"], reason="LibreOffice / PowerPoint tidak tersedia")
def test_pptx_to_pdf_real_conversion(client):
    n = upload(client, "slide.pptx", pptx_bytes())
    r = process(client, "pptx-to-pdf", [n])
    assert r.json()["results"][0]["download_name"] == "doc0.pdf"
    with pymupdf.open(stream=download(client, r).content, filetype="pdf") as doc:
        assert doc.page_count == 2 and "Presentasi Skripsi" in doc[0].get_text()


def test_limits_report_new_capabilities(client):
    body = client.get("/api/v1/limits").json()
    assert {"pptx_to_pdf_available", "docx_to_pdf_available", "ocr_available", "max_ocr_pages"} <= set(body)
